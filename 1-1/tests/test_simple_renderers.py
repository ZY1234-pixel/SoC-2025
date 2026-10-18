from pathlib import Path

from docflow.model.stages import AnalysisPage, DocumentAnalysis, Rect, SemanticElement, TypographicRole
from docflow.renderer.simple_renderers import HtmlRenderer, PlainTextRenderer, RtfRenderer, XlsxRenderer
from Code.utils import parse_formats


def _analysis():
    role = TypographicRole("body", "宋体", "Times New Roman", 10.5, 1.0)
    elements = (
        SemanticElement("title", "heading", Rect(0, 0, 100, 20), 1, ("r1",), text="标题", role_id="body"),
        SemanticElement("body", "paragraph_group", Rect(0, 20, 100, 40), 2, ("r2",), text="正文", role_id="body", payload={"image_base64": "AAAA"}),
        SemanticElement(
            "table",
            "table_group",
            Rect(0, 40, 100, 80),
            3,
            ("r3",),
            payload={"html": "<table><tr><th>列1</th><th>列2</th></tr><tr><td>A<br>下一行</td><td>B</td></tr></table>"},
            role_id="body",
        ),
    )
    return DocumentAnalysis((AnalysisPage(0, 100, 100, elements),), (role,))


def test_plain_text_renderer_includes_tables(tmp_path: Path) -> None:
    output = tmp_path / "result.txt"
    PlainTextRenderer().render(_analysis(), str(output))
    text = output.read_text(encoding="utf-8")
    assert "标题" in text and "正文" in text
    assert "列1\t列2" in text and "A\n下一行\tB" in text


def test_html_and_rtf_renderers_emit_content(tmp_path: Path) -> None:
    html_path = tmp_path / "result.html"
    rtf_path = tmp_path / "result.rtf"
    HtmlRenderer().render(_analysis(), str(html_path))
    RtfRenderer().render(_analysis(), str(rtf_path))
    html_text = html_path.read_text(encoding="utf-8")
    assert "列1" in html_text
    assert "data:image/png;base64,AAAA" not in html_text
    assert "\\rtf1" in rtf_path.read_text(encoding="ascii")
    assert "\\u21015?" in rtf_path.read_text(encoding="ascii")
    assert "\\trowd" in rtf_path.read_text(encoding="ascii")
    assert "\\cell" in rtf_path.read_text(encoding="ascii")


def test_parse_formats_accepts_new_text_formats() -> None:
    assert parse_formats("txt,html,rtf,xlsx") == ["txt", "html", "rtf", "xlsx"]


def test_xlsx_renderer_exports_tables(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    output = tmp_path / "result.xlsx"
    XlsxRenderer().render(_analysis(), str(output))
    workbook = load_workbook(output)
    assert "正文" in workbook.sheetnames and "表1" in workbook.sheetnames
    assert workbook["表1"]["A1"].value == "列1"
    assert workbook["表1"]["A2"].value == "A\n下一行"
