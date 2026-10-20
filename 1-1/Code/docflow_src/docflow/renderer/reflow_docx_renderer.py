"""Mechanical DOCX emission from a complete ReflowLayoutPlan."""

from __future__ import annotations

import base64
import fitz
import io
import math
import re
from dataclasses import replace
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image

from docflow.model.stages import (
    INLINE_FORMULA_MARKER_PATTERN,
    FlowKind,
    Rect,
    ReflowLayoutPlan,
    TextParagraphLayout,
)
from docflow.planning.text_metrics import (
    resolve_text_layout,
    visual_text,
)
from docflow.renderer.docx_utils.html_table import (
    estimate_table_cell_lines,
    fit_table_font_size,
    get_table_cell_content_nodes,
    get_table_cell_placements,
    get_table_cell_text_lines,
    get_table_column_weights,
    get_table_dimensions,
    get_table_rows,
    has_table_cell_paragraphs,
    table_cell_available_width,
    table_cell_padding_ratios,
)
from docflow.renderer.docx_utils.omml import append_latex_omml
from docflow.renderer.formula_mathjax import MathJaxFormulaRenderer
from docflow.renderer.docx_utils.table_fmt import (
    clear_table_borders,
    set_cell_margins,
    set_horizontal_table_borders,
    set_table_col_widths,
)
_ALIGNMENT = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

class ReflowDocxRenderer:
    def render(self, plan: ReflowLayoutPlan, output_path: str) -> None:
        document = self.build(plan)
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_name(f".{output.name}.tmp")
        document.save(temporary)
        temporary.replace(output)

    def build(self, plan: ReflowLayoutPlan):
        document = Document()
        roles = {role.role_id: role for role in plan.roles}
        for page_number, page in enumerate(plan.pages):
            section = document.sections[0] if page_number == 0 else document.add_section(WD_SECTION.NEW_PAGE)
            if page_number > 0:
                self._collapse_section_break(document.paragraphs[-1])
            self._set_page_geometry(section, page.geometry)
            elements = {element.element_id: element for element in page.elements}
            self._set_page_number_start(section, self._page_number_start(page, elements))
            edge_ids = {
                element.element_id
                for element in page.elements
                if self._is_edge_vertical_text(element, page.geometry)
            }
            header_ids = tuple(identifier for identifier in page.header_element_ids if identifier not in edge_ids)
            footer_ids = tuple(identifier for identifier in page.footer_element_ids if identifier not in edge_ids)
            self._set_furniture_distances(section, page, elements, header_ids, footer_ids)
            usable_width = self._usable_width(page.geometry)
            if page.header_element_ids:
                section.header.is_linked_to_previous = False
                if header_ids:
                    self._render_furniture(section.header, header_ids, elements, roles, page.fit_scale, usable_width)
                else:
                    self._clear_container(section.header)
            if page.footer_element_ids:
                section.footer.is_linked_to_previous = False
                if footer_ids:
                    self._render_furniture(section.footer, footer_ids, elements, roles, page.fit_scale, usable_width)
                else:
                    self._clear_container(section.footer)
            body = document
            for identifier in page.header_element_ids + page.footer_element_ids:
                if identifier in edge_ids:
                    self._render_anchored_edge_text(body, elements[identifier], roles, page.fit_scale)
            for flow in page.sections:
                if flow.kind == FlowKind.SINGLE:
                    for identifier in flow.element_ids:
                        self._render_element(body, elements[identifier], roles, page.fit_scale, usable_width)
                elif flow.kind == FlowKind.WRAPPED:
                    self._render_wrapped(body, flow, elements, roles, page.fit_scale, usable_width)
                elif flow.kind == FlowKind.SEQUENTIAL_COLUMNS:
                    self._render_columns(body, flow, elements, roles, page.fit_scale)
                else:
                    self._render_grid(body, flow, elements, roles, page.fit_scale, usable_width)
        if plan.pages:
            self._collapse_section_break(document.add_paragraph())
        return document

    @staticmethod
    def _collapse_section_break(paragraph) -> None:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(1)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        properties = paragraph._p.get_or_add_pPr()
        spacing = properties.find(qn("w:spacing"))
        spacing.set(qn("w:line"), "1")
        mark_properties = properties.find(qn("w:rPr"))
        if mark_properties is None:
            mark_properties = OxmlElement("w:rPr")
            properties.append(mark_properties)
        mark_properties.append(OxmlElement("w:vanish"))
        if not paragraph.runs:
            paragraph.add_run()
        for run in paragraph.runs:
            run.font.size = Pt(0.5)
            run._r.get_or_add_rPr().append(OxmlElement("w:vanish"))

    @staticmethod
    def _set_page_geometry(section, geometry) -> None:
        for grid in section._sectPr.xpath("./w:docGrid"):
            section._sectPr.remove(grid)
        section.page_width = Pt(geometry.width_pt)
        section.page_height = Pt(geometry.height_pt)
        section.top_margin = Pt(geometry.margin_top_pt)
        section.right_margin = Pt(geometry.margin_right_pt)
        section.bottom_margin = Pt(geometry.margin_bottom_pt)
        section.left_margin = Pt(geometry.margin_left_pt)

    @staticmethod
    def _page_number_start(page, elements) -> int | None:
        for identifier in page.header_element_ids + page.footer_element_ids:
            element = elements[identifier]
            if not ReflowDocxRenderer._is_page_number_element(element):
                continue
            match = re.search(r"\d{1,4}", str(element.text or ""))
            if match:
                return int(match.group(0))
        return None

    @staticmethod
    def _set_page_number_start(section, start: int | None) -> None:
        properties = section._sectPr
        for node in properties.xpath("./w:pgNumType"):
            properties.remove(node)
        if start is None:
            return
        page_number_type = OxmlElement("w:pgNumType")
        page_number_type.set(qn("w:start"), str(start))
        properties.append(page_number_type)

    def _render_furniture(self, container, identifiers, elements, roles, fit_scale, usable_width) -> None:
        self._clear_container(container)
        if not identifiers:
            return
        textual = [
            element
            for element in (elements[identifier] for identifier in identifiers)
            if str(element.text or "").strip()
        ]
        if textual and len(textual) < len(identifiers) and all(
            elements[identifier].kind == "header" for identifier in identifiers
        ):
            self._render_header_line(container, identifiers, elements, roles, fit_scale, usable_width)
            return
        if textual and len(textual) == len(identifiers) and all(element.kind == "header" for element in textual):
            ordered = sorted(
                textual,
                key=lambda element: float(element.payload.get("source_bbox", (0, 0, 0, 0))[1]),
            )
            centers = [
                (float(element.payload.get("source_bbox", (0, 0, 0, 0))[1])
                 + float(element.payload.get("source_bbox", (0, 0, 0, 0))[3]))
                / 2.0
                for element in ordered
            ]
            heights = [
                float(element.payload.get("source_bbox", (0, 0, 0, 0))[3])
                - float(element.payload.get("source_bbox", (0, 0, 0, 0))[1])
                for element in ordered
            ]
            has_span_rows = any(len(row.spans) > 1 for element in ordered for row in element.text_rows)
            if len(ordered) == 1 or has_span_rows or max(centers) - min(centers) <= max(heights, default=1.0):
                self._render_header_texts(container, ordered, roles, fit_scale, usable_width)
                return
        if textual and len(textual) == len(identifiers) and any(
            element.kind in {"footer", "page_number"} for element in textual
        ):
            ordered = sorted(
                textual,
                key=lambda element: float(element.payload.get("source_bbox", (0, 0, 0, 0))[0]),
            )
            self._render_footer_texts(container, ordered, roles, fit_scale, usable_width)
            return
        if not textual and all(elements[identifier].payload.get("image_base64") for identifier in identifiers):
            for identifier in identifiers:
                element = elements[identifier]
                data = self._decode_image(element.payload.get("image_base64"))
                if not data:
                    continue
                bbox = element.payload.get("source_bbox") or (0, 0, 1, 1)
                source_scale = float(element.payload.get("source_scale", 1.0))
                paragraph = container.add_paragraph()
                picture = paragraph.add_run().add_picture(
                    io.BytesIO(data),
                    width=Pt(max((float(bbox[2]) - float(bbox[0])) * source_scale, 0.5)),
                )
                if element.kind == "header":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    self._anchor_picture(
                        picture._inline,
                        float(bbox[0]) * source_scale,
                        float(bbox[1]) * source_scale,
                    )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(1)
                paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            return
        if len(identifiers) == 1:
            element = elements[identifiers[0]]
            if element.text_structure.orientation == "vertical":
                self._write_vertical_text(container, element, roles, fit_scale)
                return
            paragraph = container.add_paragraph()
            self._write_text(paragraph, element, roles, fit_scale)
            self._apply_payload_text_color(paragraph, element)
            return
        table = container.add_table(rows=1, cols=3, width=Pt(usable_width))
        self._format_layout_table(table, (usable_width / 3,) * 3)
        for cell in table.rows[0].cells:
            self._clear_container(cell)
        ordered_identifiers = sorted(
            identifiers,
            key=lambda identifier: (
                min(
                    int(
                        (
                            float(elements[identifier].payload.get("source_bbox", (0, 0, 0, 0))[0])
                            + float(elements[identifier].payload.get("source_bbox", (0, 0, 0, 0))[2])
                        )
                        / 2.0
                        / max(float(elements[identifier].payload.get("page_width_px", 1.0)), 1.0)
                        * 3
                    ),
                    2,
                ),
                float(elements[identifier].payload.get("source_bbox", (0, 0, 0, 0))[1]),
            ),
        )
        for identifier in ordered_identifiers:
            element = elements[identifier]
            bbox = element.payload.get("source_bbox") or (0, 0, 1, 1)
            page_width_px = max(float(element.payload.get("page_width_px", 1.0)), 1.0)
            column = min(int(((float(bbox[0]) + float(bbox[2])) / 2.0) / page_width_px * 3), 2)
            cell = table.cell(0, column)
            paragraph = cell.add_paragraph()
            paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT)[column]
            if element.text_structure.orientation == "vertical":
                self._write_vertical_text(cell, element, roles, fit_scale)
                continue
            if element.payload.get("image_base64") and not str(element.text or "").strip():
                data = self._decode_image(element.payload.get("image_base64"))
                if data:
                    width = (float(bbox[2]) - float(bbox[0])) / page_width_px * usable_width * fit_scale
                    paragraph.add_run().add_picture(io.BytesIO(data), width=Pt(max(width, 0.5)))
                continue
            self._write_text(paragraph, element, roles, fit_scale)
            self._apply_payload_text_color(paragraph, element)
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.right_indent = Pt(0)
            paragraph.alignment = (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT)[column]

    def _render_header_line(self, container, identifiers, elements, roles, fit_scale: float, usable_width: float) -> None:
        ordered = sorted(
            (elements[identifier] for identifier in identifiers),
            key=lambda element: float(element.payload.get("source_bbox", (0, 0, 0, 0))[0]),
        )
        paragraph = container.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        image_height = 0.0
        for element in ordered:
            data = self._decode_image(element.payload.get("image_base64"))
            if data:
                target_width = float(element.payload.get("target_width_pt") or 0.0)
                target_height = float(element.payload.get("target_height_pt") or 0.0)
                image_width, height = self._image_size(
                    data,
                    target_width or 1.0,
                    target_height if target_height > 0 else float("inf"),
                )
                paragraph.add_run().add_picture(
                    io.BytesIO(data),
                    width=Pt(image_width),
                    height=Pt(height),
                )
                image_height = max(image_height, height)
                continue
            if str(element.text or "").strip():
                self._write_text(paragraph, element, roles, fit_scale, usable_width)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        line_height = max((run.font.size.pt * 1.2 for run in paragraph.runs if run.font.size), default=10.5)
        paragraph.paragraph_format.line_spacing = Pt(max(image_height, line_height) * 1.05)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST

    def _render_columns(self, container, flow, elements, roles, fit_scale) -> None:
        column_section = container.add_section(WD_SECTION.CONTINUOUS)
        self._set_columns(column_section, flow.column_widths_pt, flow.gutter_pt)
        column = 0
        started_columns = set()
        for identifier in flow.element_ids:
            element = elements[identifier]
            if element.kind == "figure_group":
                element = replace(
                    element,
                    payload={**element.payload, "alignment": "center"},
                )
            element_column = int(element.payload.get("column", 0))
            while element_column > column:
                self._write_column_break(container)
                column += 1
            offset = float(element.payload.get("column_offset_pt") or 0.0)
            if column not in started_columns and offset > 0.0:
                spacer = container.add_paragraph()
                spacer.paragraph_format.space_before = Pt(0)
                spacer.paragraph_format.space_after = Pt(0)
                spacer.paragraph_format.line_spacing = Pt(offset * fit_scale)
                spacer.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                run = spacer.add_run("")
                run.font.size = Pt(1)
            started_columns.add(column)
            self._render_element(
                container,
                element,
                roles,
                fit_scale,
                flow.column_widths_pt[element_column],
            )
        restore_section = container.add_section(WD_SECTION.CONTINUOUS)
        usable_width = sum(flow.column_widths_pt) + flow.gutter_pt * (len(flow.column_widths_pt) - 1)
        self._set_columns(restore_section, (usable_width,), 0.0)

    @staticmethod
    def _set_columns(section, widths, gutter_pt: float) -> None:
        properties = section._sectPr
        for node in properties.xpath("./w:cols"):
            properties.remove(node)
        columns = OxmlElement("w:cols")
        if len(widths) == 1:
            columns.set(qn("w:space"), "0")
        else:
            columns.set(qn("w:num"), str(len(widths)))
            columns.set(qn("w:space"), str(round(max(gutter_pt, 0.0) * 20)))
            columns.set(qn("w:equalWidth"), "0")
            for index, width in enumerate(widths):
                column = OxmlElement("w:col")
                column.set(qn("w:w"), str(round(width * 20)))
                if index + 1 < len(widths):
                    column.set(qn("w:space"), str(round(max(gutter_pt, 0.0) * 20)))
                columns.append(column)
        properties.append(columns)

    @staticmethod
    def _write_column_break(container) -> None:
        paragraph = container.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(1)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.add_run().add_break(WD_BREAK.COLUMN)

    def _render_wrapped(self, container, flow, elements, roles, fit_scale, container_width) -> None:
        floating = elements[flow.floating_element_id]
        width = flow.floating_width_pt * fit_scale
        table = container.add_table(rows=1, cols=1)
        self._format_layout_table(table, (width,))
        self._float_table(
            table,
            flow.floating_side,
            flow.floating_offset_x_pt * fit_scale,
            flow.floating_offset_y_pt * fit_scale,
            flow.gutter_pt * fit_scale,
        )
        cell = table.cell(0, 0)
        set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
        self._clear_container(cell)
        data = self._decode_image(floating.payload.get("image_base64"))
        if floating.kind == "figure_group" and data:
            paragraph = cell.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.add_run().add_picture(io.BytesIO(data), width=Pt(max(width, 0.5)))
            self._write_caption(
                cell,
                floating.payload.get("caption"),
                roles,
                fit_scale,
                floating.payload.get("caption_alignment"),
                roles.get(floating.role_id),
                floating.payload.get("caption_font_size_pt"),
            )
        else:
            self._render_element(cell, floating, roles, fit_scale, width)
        self._collapse_trailing_paragraph(cell)
        self._collapse_trailing_paragraph(container)
        for identifier in flow.element_ids:
            if identifier != flow.floating_element_id:
                self._render_element(container, elements[identifier], roles, fit_scale, container_width)

    @staticmethod
    def _float_table(table, side: str, offset_x_pt: float, offset_y_pt: float, gutter_pt: float) -> None:
        properties = table._tbl.tblPr
        positioning = OxmlElement("w:tblpPr")
        distance = str(round(max(gutter_pt / 2.0, 2.0) * 20))
        positioning.set(qn("w:leftFromText"), distance)
        positioning.set(qn("w:rightFromText"), distance)
        positioning.set(qn("w:topFromText"), "0")
        positioning.set(qn("w:bottomFromText"), "0")
        positioning.set(qn("w:vertAnchor"), "text")
        positioning.set(qn("w:horzAnchor"), "text")
        if side == "left":
            positioning.set(qn("w:tblpX"), str(round(offset_x_pt * 20)))
        else:
            positioning.set(qn("w:tblpXSpec"), side)
        positioning.set(qn("w:tblpY"), str(round(offset_y_pt * 20)))
        properties.append(positioning)

    @staticmethod
    def _anchor_picture(inline, x_pt: float, y_pt: float) -> None:
        inline.tag = qn("wp:anchor")
        for name, value in {
            "distT": "0",
            "distB": "0",
            "distL": "0",
            "distR": "0",
            "simplePos": "0",
            "relativeHeight": "0",
            "behindDoc": "0",
            "locked": "0",
            "layoutInCell": "1",
            "allowOverlap": "1",
        }.items():
            inline.set(name, value)
        simple_position = OxmlElement("wp:simplePos")
        simple_position.set("x", "0")
        simple_position.set("y", "0")
        inline.insert(0, simple_position)
        for index, (axis, offset) in enumerate((("H", x_pt), ("V", y_pt)), start=1):
            position = OxmlElement(f"wp:position{axis}")
            position.set("relativeFrom", "page")
            value = OxmlElement("wp:posOffset")
            value.text = str(round(offset * 12700))
            position.append(value)
            inline.insert(index, position)
        inline.insert(4, OxmlElement("wp:wrapNone"))

    def _render_grid(self, container, flow, elements, roles, fit_scale, usable_width: float | None = None) -> None:
        usable_width = usable_width or sum(flow.column_widths_pt) + flow.gutter_pt
        if self._render_heading_icon_table(container, flow, elements, roles, fit_scale, usable_width):
            return
        row_count = max(cell.row + cell.row_span for cell in flow.grid_cells)
        column_count = len(flow.column_widths_pt)
        table = container.add_table(rows=row_count, cols=column_count)
        self._format_layout_table(table, flow.column_widths_pt, flow.gutter_pt)
        fixed_rows = {
            row
            for cell in flow.grid_cells
            if cell.row_span > 1
            and len(cell.element_ids) == 1
            and elements[cell.element_ids[0]].kind == "figure_group"
            for row in range(cell.row, cell.row + cell.row_span)
        }
        for index, row in enumerate(table.rows):
            if index >= len(flow.row_heights_pt):
                break
            row.height = Pt(flow.row_heights_pt[index] * fit_scale)
            row.height_rule = (
                WD_ROW_HEIGHT_RULE.EXACTLY
                if index in fixed_rows
                else WD_ROW_HEIGHT_RULE.AT_LEAST
            )
        target_cells = {}
        for grid_cell in flow.grid_cells:
            cell = table.cell(grid_cell.row, grid_cell.column)
            floats_across_rows = (
                grid_cell.row_span > 1
                and len(grid_cell.element_ids) == 1
                and elements[grid_cell.element_ids[0]].kind == "figure_group"
            )
            if grid_cell.row_span > 1 or grid_cell.column_span > 1:
                end_row = grid_cell.row if floats_across_rows else grid_cell.row + grid_cell.row_span - 1
                cell = cell.merge(table.cell(end_row, grid_cell.column + grid_cell.column_span - 1))
            half_gutter_twips = int(max(flow.gutter_pt, 0.0) * 10)
            set_cell_margins(
                cell,
                top=0,
                bottom=0,
                start=half_gutter_twips if grid_cell.column > 0 else 0,
                end=half_gutter_twips if grid_cell.column + grid_cell.column_span < column_count else 0,
            )
            target_cells[(grid_cell.row, grid_cell.column)] = cell
        for properties in table._tbl.xpath(".//w:tcPr"):
            if properties.find(qn("w:hideMark")) is None:
                properties.append(OxmlElement("w:hideMark"))
        for row in table.rows:
            for cell in row.cells:
                if cell._tc.xpath("./w:p | ./w:tbl"):
                    self._clear_container(cell)
        for grid_cell in flow.grid_cells:
            cell = target_cells[(grid_cell.row, grid_cell.column)]
            container_width = (
                sum(flow.column_widths_pt[grid_cell.column : grid_cell.column + grid_cell.column_span])
                + flow.gutter_pt * (grid_cell.column_span - 1)
            )
            floats_across_rows = (
                grid_cell.row_span > 1
                and len(grid_cell.element_ids) == 1
                and elements[grid_cell.element_ids[0]].kind == "figure_group"
            )
            for identifier in grid_cell.element_ids:
                element = elements[identifier]
                if element.kind == "figure_group" and floats_across_rows:
                    element = replace(element, payload={**element.payload, "float_in_grid": True})
                self._render_element(cell, element, roles, fit_scale, container_width)
        for row in table.rows:
            for cell in row.cells:
                if not cell.paragraphs and not cell.tables:
                    cell.add_paragraph()
                    self._collapse_trailing_paragraph(cell)
        self._collapse_trailing_paragraph(container)

    def _render_heading_icon_table(
        self,
        container,
        flow,
        elements,
        roles,
        fit_scale: float,
        usable_width: float | None = None,
    ) -> bool:
        usable_width = usable_width or sum(flow.column_widths_pt) + flow.gutter_pt
        if len(flow.grid_cells) != 2 or any(cell.row != 0 or cell.row_span != 1 for cell in flow.grid_cells):
            return False
        cells = sorted(flow.grid_cells, key=lambda cell: cell.column)
        if [cell.column for cell in cells] != [0, 1] or any(len(cell.element_ids) != 1 for cell in cells):
            return False
        first = elements[cells[0].element_ids[0]]
        second = elements[cells[1].element_ids[0]]
        if first.kind == "figure_group" and second.kind == "heading":
            figure, heading = first, second
        elif first.kind == "heading" and second.kind == "figure_group":
            heading, figure = first, second
        else:
            return False
        figure_bbox = figure.payload.get("source_bbox") or (0, 0, 1, 1)
        heading_bbox = heading.payload.get("source_bbox") or (0, 0, 1, 1)
        overlap = max(0.0, min(figure_bbox[3], heading_bbox[3]) - max(figure_bbox[1], heading_bbox[1]))
        if overlap / max(min(figure_bbox[3] - figure_bbox[1], heading_bbox[3] - heading_bbox[1]), 1.0) < 0.50:
            return False
        data = self._decode_image(figure.payload.get("image_base64"))
        if not data:
            return False
        target_width = float(figure.payload.get("target_width_pt") or 0.0)
        target_height = float(figure.payload.get("target_height_pt") or 0.0)
        source_width = max(float(figure_bbox[2]) - float(figure_bbox[0]), 1.0)
        image_width, image_height = self._image_size(
            data,
            target_width or source_width,
            target_height if target_height > 0 else float("inf"),
        )
        heading_line_height = heading.text_layout[0].line_height_pt if heading.text_layout else 12.0
        heading_width = float(heading.payload.get("target_width_pt") or max(heading_bbox[2] - heading_bbox[0], 1.0))
        group_width = image_width + flow.gutter_pt + heading_width
        left_space = max((usable_width - group_width) / 2.0, 0.0)
        right_space = max(usable_width - left_space - group_width, 0.0)
        table = container.add_table(rows=1, cols=4)
        self._format_layout_table(table, (left_space, image_width, heading_width, right_space))
        table_width_element = table._tbl.tblPr.find(qn("w:tblW"))
        table_width_element.set(qn("w:type"), "pct")
        table_width_element.set(qn("w:w"), "5000")
        table.rows[0].height = Pt(max(image_height, heading_line_height))
        table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        figure_cell = table.cell(0, 1)
        heading_cell = table.cell(0, 2)
        set_cell_margins(
            figure_cell,
            top=0,
            bottom=0,
            start=0,
            end=int(max(flow.gutter_pt, 0.0) * 10),
        )
        set_cell_margins(
            heading_cell,
            top=0,
            bottom=0,
            start=int(max(flow.gutter_pt, 0.0) * 10),
            end=0,
        )
        figure_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        heading_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        figure_paragraph = figure_cell.paragraphs[0]
        figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        figure_paragraph.paragraph_format.space_before = Pt(0)
        figure_paragraph.paragraph_format.space_after = Pt(0)
        figure_paragraph.add_run().add_picture(
            io.BytesIO(data),
            width=Pt(image_width),
            height=Pt(image_height),
        )
        heading_paragraph = heading_cell.paragraphs[0]
        self._write_text(heading_paragraph, heading, roles, fit_scale, heading_width)
        heading_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading_paragraph.paragraph_format.space_before = Pt(0)
        heading_paragraph.paragraph_format.space_after = Pt(0)
        heading_paragraph.paragraph_format.line_spacing = Pt(heading_line_height)
        heading_paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        self._collapse_trailing_paragraph(container)
        return True

    def _render_element(self, container, element, roles, fit_scale: float, container_width: float) -> None:
        if element.kind in {"heading", "paragraph_group", "caption"}:
            if element.text_structure.orientation == "vertical":
                self._write_vertical_text(container, element, roles, fit_scale)
                return
            role = roles.get(element.role_id) or self._body_role(roles)
            layouts = element.text_layout or resolve_text_layout(element, role, container_width, fit_scale)
            for layout in layouts:
                self._write_resolved_text(container.add_paragraph(), element, role, layout)
            return
        if element.kind != "equation_group" or element.payload.get("number"):
            self._write_block_spacing(container, element, fit_scale)
        if element.kind == "figure_group":
            if element.payload.get("caption_position") == "before":
                self._write_caption(container, element.payload.get("caption"), roles, fit_scale, element.payload.get("caption_alignment"), roles.get(element.role_id), element.payload.get("caption_font_size_pt"), container_width)
            self._write_image(container, element, fit_scale, container_width)
            if element.payload.get("caption_position") != "before":
                self._write_caption(container, element.payload.get("caption"), roles, fit_scale, element.payload.get("caption_alignment"), roles.get(element.role_id), element.payload.get("caption_font_size_pt"), container_width)
            return
        if element.kind == "equation_group":
            self._write_equation(container, element, roles, fit_scale, container_width)
            return
        if element.kind == "table_group":
            has_native_table = bool(element.payload.get("html"))
            if element.payload.get("caption_position") == "before":
                self._write_caption(container, element.payload.get("caption"), roles, fit_scale, element.payload.get("caption_alignment"), roles.get(element.role_id), element.payload.get("caption_font_size_pt"), container_width)
            if has_native_table:
                self._write_native_table(container, element, roles, fit_scale, container_width)
            else:
                self._write_image(container, element, fit_scale, container_width)
            if element.payload.get("caption_position") != "before":
                self._write_caption(container, element.payload.get("caption"), roles, fit_scale, element.payload.get("caption_alignment"), roles.get(element.role_id), element.payload.get("caption_font_size_pt"), container_width)

    def _write_vertical_text(self, container, element, roles, fit_scale: float) -> None:
        bbox = element.payload.get("source_bbox") or (0, 0, 1, 1)
        width = max((float(bbox[2]) - float(bbox[0])) * float(element.payload.get("source_scale", 1.0)), 1.0)
        if hasattr(container, "_tc"):
            self._shade(container._tc.get_or_add_tcPr(), element.payload.get("background_color"))
        table = container.add_table(rows=1, cols=1)
        self._format_layout_table(table, (width,))
        cell = table.cell(0, 0)
        self._shade(cell._tc.get_or_add_tcPr(), element.payload.get("background_color"))
        paragraph = cell.paragraphs[0]
        self._write_text(paragraph, element, roles, fit_scale, width)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        self._collapse_trailing_paragraph(container)

    @staticmethod
    def _is_edge_vertical_text(element, geometry) -> bool:
        if element.text_structure.orientation != "vertical" or not str(element.text or "").strip():
            return False
        bbox = element.payload.get("source_bbox") or ()
        if len(bbox) != 4:
            return False
        source_scale = float(element.payload.get("source_scale", 1.0))
        x1, y1, x2, y2 = (float(value) for value in bbox)
        width = max(x2 - x1, 1.0)
        height = max(y2 - y1, 1.0)
        return (
            width * source_scale <= geometry.width_pt * 0.09
            and height >= width * 2.5
            and (x2 * source_scale <= geometry.width_pt * 0.12 or x1 * source_scale >= geometry.width_pt * 0.88)
        )

    def _render_anchored_edge_text(self, container, element, roles, fit_scale: float) -> None:
        bbox = element.payload.get("source_bbox") or (0, 0, 1, 1)
        source_scale = float(element.payload.get("source_scale", 1.0))
        x1, y1, x2, _y2 = (float(value) for value in bbox)
        width = max((x2 - x1) * source_scale, 1.0)
        table = container.add_table(rows=1, cols=1)
        self._format_layout_table(table, (width,))
        cell = table.cell(0, 0)
        set_cell_margins(cell, top=0, bottom=0, start=0, end=0)
        self._clear_container(cell)
        paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        self._write_text(paragraph, element, roles, fit_scale, width)
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.right_indent = Pt(0)
        self._position_page_table(table, x1 * source_scale, y1 * source_scale)

    @staticmethod
    def _position_page_table(table, x_pt: float, y_pt: float) -> None:
        properties = table._tbl.tblPr
        positioning = OxmlElement("w:tblpPr")
        positioning.set(qn("w:leftFromText"), "0")
        positioning.set(qn("w:rightFromText"), "0")
        positioning.set(qn("w:topFromText"), "0")
        positioning.set(qn("w:bottomFromText"), "0")
        positioning.set(qn("w:vertAnchor"), "page")
        positioning.set(qn("w:horzAnchor"), "page")
        positioning.set(qn("w:tblpX"), str(round(x_pt * 20)))
        positioning.set(qn("w:tblpY"), str(round(y_pt * 20)))
        properties.append(positioning)
        overlap = OxmlElement("w:tblOverlap")
        overlap.set(qn("w:val"), "overlap")
        properties.append(overlap)

    def _write_text(self, paragraph, element, roles, fit_scale: float, container_width: float | None = None) -> None:
        role = roles.get(element.role_id) or self._body_role(roles)
        if role is None:
            paragraph.add_run(element.text)
            return
        width = container_width
        if width is None:
            bbox = self._content_bbox(element)
            width = max(bbox.width * float(element.payload.get("source_scale", 1.0)), 1.0)
        layouts = element.text_layout or resolve_text_layout(element, role, width, fit_scale)
        if layouts:
            layout = layouts[0]
            if self._is_page_number_element(element):
                self._write_page_number_text(paragraph, element, role, layout)
            else:
                self._write_resolved_text(paragraph, element, role, layout)

    @staticmethod
    def _is_page_number_element(element) -> bool:
        text = str(element.text or "").strip()
        return (
            element.kind == "page_number"
            or (element.kind in {"header", "footer"} and bool(re.fullmatch(r"\d{1,4}", text)))
            or (element.kind == "footer" and bool(re.match(r"^\s*\d{1,4}\s+\S", text)))
        )

    def _write_page_number_text(self, paragraph, element, role, layout) -> None:
        layout = self._page_number_layout(element, layout)
        paragraph.paragraph_format.space_before = Pt(layout.space_before_pt)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(layout.left_indent_pt)
        paragraph.paragraph_format.right_indent = Pt(layout.right_indent_pt)
        paragraph.alignment = _ALIGNMENT[layout.alignment]
        paragraph.paragraph_format.line_spacing = Pt(layout.line_height_pt)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        self._append_page_number_runs(paragraph, element, role, layout)

    @staticmethod
    def _page_number_layout(element, layout):
        bbox = Rect.from_sequence(element.payload.get("source_bbox") or (0, 0, 1, 1))
        page_width = max(float(element.payload.get("page_width_px") or 1.0), 1.0)
        return replace(
            layout,
            alignment=ReflowDocxRenderer._infer_furniture_alignment(bbox, 0.0, page_width),
            first_line_indent_pt=0.0,
            left_indent_pt=0.0,
            right_indent_pt=0.0,
        )

    def _append_page_number_runs(self, paragraph, element, role, layout) -> None:
        text = str(element.text or "")
        matches = list(re.finditer(r"\d{1,4}", text))
        if not matches:
            self._write_text_run(paragraph, text, role, layout)
            return
        if element.kind == "footer":
            match = matches[0]
            self._write_text_run(paragraph, text[: match.start()], role, layout)
            self._append_field_run(paragraph, role, layout, "PAGE", match.group(0))
            self._write_text_run(paragraph, text[match.end() :], role, layout)
            return
        position = 0
        for index, match in enumerate(matches):
            self._write_text_run(paragraph, text[position : match.start()], role, layout)
            field = "PAGE" if index == 0 else "NUMPAGES"
            self._append_field_run(paragraph, role, layout, field, match.group(0))
            position = match.end()
        self._write_text_run(paragraph, text[position:], role, layout)

    def _append_field_run(self, paragraph, role, layout, field: str, result: str) -> None:
        run = paragraph.add_run()
        self._style_run(run, role, 1.0, font_size_pt=layout.font_size_pt)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = f" {field} "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        run._r.extend((begin, instruction, separate))
        self._write_text_run(paragraph, result, role, layout)
        end_run = paragraph.add_run()
        self._style_run(end_run, role, 1.0, font_size_pt=layout.font_size_pt)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        end_run._r.append(end)

    def _write_resolved_text(self, paragraph, element, role, layout) -> None:
        background_color = element.payload.get("background_color")
        paragraph.paragraph_format.space_before = Pt(layout.space_before_pt)
        paragraph.paragraph_format.space_after = Pt(0)
        centered = layout.alignment == "center"
        asymmetric_center_indent = centered and abs(layout.left_indent_pt - layout.right_indent_pt) > 1.0
        paragraph.paragraph_format.first_line_indent = Pt(0 if centered else layout.first_line_indent_pt)
        paragraph.paragraph_format.left_indent = Pt(0 if asymmetric_center_indent else layout.left_indent_pt)
        paragraph.paragraph_format.right_indent = Pt(layout.right_indent_pt)
        paragraph.alignment = _ALIGNMENT[layout.alignment]
        if background_color and element.kind == "heading":
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.left_indent = Pt(0)
            paragraph.paragraph_format.right_indent = Pt(0)
        if layout.right_tab_stop_pt is not None:
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Pt(layout.right_tab_stop_pt),
                WD_TAB_ALIGNMENT.RIGHT,
            )
        paragraph.paragraph_format.line_spacing = Pt(layout.line_height_pt)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        formulas = element.payload.get("inline_formulas") or ()
        position = 0
        for match in INLINE_FORMULA_MARKER_PATTERN.finditer(layout.text):
            if match.start() > position:
                self._write_text_run(paragraph, layout.text[position:match.start()], role, layout)
            index = int(match.group(1))
            formula = formulas[index] if index < len(formulas) else {}
            latex = str(formula.get("latex") or match.group(0)).strip("$")
            if not append_latex_omml(paragraph, latex, layout.font_size_pt):
                self._write_text_run(paragraph, f"${latex}$", role, layout)
            position = match.end()
        if position < len(layout.text):
            self._write_text_run(paragraph, layout.text[position:], role, layout)
        if background_color and element.kind == "heading":
            for run in paragraph.runs:
                if run.text:
                    run.text = f"\u00a0{run.text}\u00a0"
                self._shade(run._element.get_or_add_rPr(), background_color)
        else:
            self._shade(paragraph._p.get_or_add_pPr(), background_color)
        paragraph.paragraph_format.widow_control = False
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.keep_with_next = False

    @staticmethod
    def _write_text_run(paragraph, text: str, role, layout) -> None:
        if not text:
            return
        run = paragraph.add_run(text)
        ReflowDocxRenderer._style_run(run, role, 1.0, font_size_pt=layout.font_size_pt)
        if layout.font_width_scale_pct != 100:
            width = OxmlElement("w:w")
            width.set(qn("w:val"), str(layout.font_width_scale_pct))
            run._element.get_or_add_rPr().append(width)
        if layout.character_spacing_pt:
            spacing = OxmlElement("w:spacing")
            spacing.set(qn("w:val"), str(round(layout.character_spacing_pt * 20)))
            run._element.get_or_add_rPr().append(spacing)

    @staticmethod
    def _set_furniture_distances(section, page, elements, header_ids, footer_ids) -> None:
        if header_ids:
            section.header_distance = Pt(
                min(float(elements[identifier].payload["source_bbox"][1]) for identifier in header_ids)
                * float(elements[header_ids[0]].payload.get("source_scale", 1.0))
            )
        else:
            section.header_distance = Pt(0)
        if footer_ids:
            reference = elements[footer_ids[0]]
            page_height = float(reference.payload.get("page_height_px", 1.0))
            bottom = max(float(elements[identifier].payload["source_bbox"][3]) for identifier in footer_ids)
            section.footer_distance = Pt((page_height - bottom) * float(reference.payload.get("source_scale", 1.0)))
        else:
            section.footer_distance = Pt(0)

    @staticmethod
    def _write_block_spacing(container, element, fit_scale: float) -> None:
        spacing = float(element.payload.get("space_before_pt", 0.0)) * fit_scale
        if spacing <= 0:
            return
        paragraph = container.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(max(spacing - 1.0, 0.0))
        paragraph.paragraph_format.line_spacing = Pt(1)

    def _write_caption(self, container, text, roles, fit_scale: float, alignment=None, fallback_role=None, font_size_pt=None, container_width=None) -> None:
        if not text:
            return
        role = fallback_role or next((role for role_id, role in roles.items() if role_id.startswith("caption_")), None) or self._body_role(roles)
        paragraph = container.add_paragraph()
        paragraph.alignment = _ALIGNMENT.get(alignment, WD_ALIGN_PARAGRAPH.CENTER)
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(0)
        if "\t" in str(text) and container_width:
            paragraph.paragraph_format.tab_stops.add_tab_stop(Pt(container_width / 2.0), WD_TAB_ALIGNMENT.CENTER)
        run = paragraph.add_run(str(text))
        font_size = max(round(float(font_size_pt) * fit_scale * 2) / 2.0, 0.5) if font_size_pt else None
        self._style_run(run, role, fit_scale if font_size is None else 1.0, font_size_pt=font_size)
        rendered_size = font_size or max(round(role.font_size_pt * fit_scale * 2) / 2.0, 0.5)
        paragraph.paragraph_format.line_spacing = Pt(rendered_size * 1.1)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY

    def _write_image(self, container, element, fit_scale: float, container_width: float) -> None:
        data = self._decode_image(element.payload.get("image_base64"))
        if not data:
            return
        bbox = element.payload.get("primary_bbox") or element.payload.get("source_bbox") or (0, 0, 1, 1)
        source_width = (float(bbox[2]) - float(bbox[0])) * float(element.payload.get("source_scale", 1.0))
        element_fit_scale = 1.0 if element.kind == "figure_group" else fit_scale
        target_width = float(element.payload.get("target_width_pt") or 0.0) * element_fit_scale
        target_height = float(element.payload.get("target_height_pt") or 0.0) * element_fit_scale
        max_width = min(container_width * float(element.payload.get("width_fraction", 1.0)), source_width)
        if target_width > 0:
            max_width = min(max_width, target_width)
        if element.kind != "figure_group":
            max_width *= fit_scale
        caption = str(element.payload.get("caption") or "")
        if caption and target_height > 0:
            caption_size = float(element.payload.get("caption_font_size_pt") or 10.0) * element_fit_scale
            target_height = max(target_height - len(caption.splitlines()) * caption_size * 1.1, 1.0)
        image_width, image_height = self._image_size(
            data,
            max(max_width, 0.5),
            target_height if target_height > 0 else float("inf"),
        )
        paragraph = container.add_paragraph()
        paragraph.alignment = _ALIGNMENT.get(element.payload.get("alignment"), WD_ALIGN_PARAGRAPH.CENTER)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Pt(float(element.payload.get("left_indent_pt") or 0.0))
        paragraph.paragraph_format.right_indent = Pt(float(element.payload.get("right_indent_pt") or 0.0))
        picture = paragraph.add_run().add_picture(
            io.BytesIO(data),
            width=Pt(image_width),
            height=Pt(image_height),
        )
        if element.kind == "equation_group":
            paragraph.paragraph_format.space_before = Pt(
                float(element.payload.get("space_before_pt", 0.0)) * fit_scale
            )
            paragraph.paragraph_format.line_spacing = Pt(max(image_height, 0.5))
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        if element.payload.get("float_in_grid"):
            self._float_picture(picture._inline, element.payload.get("alignment"))
            paragraph.paragraph_format.line_spacing = Pt(1)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY

    @classmethod
    def _float_picture(cls, inline, alignment) -> None:
        cls._anchor_picture(inline, 0.0, 0.0)
        horizontal = inline.find(qn("wp:positionH"))
        horizontal.set("relativeFrom", "column")
        position = horizontal[0]
        position.tag = qn("wp:align")
        position.text = alignment if alignment in {"left", "center", "right"} else "left"
        inline.find(qn("wp:positionV")).set("relativeFrom", "paragraph")

    def _write_equation(self, container, element, roles, fit_scale: float, container_width: float) -> None:
        number = str(element.payload.get("number") or "")
        latex = str(element.payload.get("latex") or "").strip()
        target_width = float(element.payload.get("target_width_pt") or 0.0) * fit_scale
        target_height = float(element.payload.get("target_height_pt") or 0.0) * fit_scale
        role = roles.get(element.role_id) or self._body_role(roles)
        formula_size = role.font_size_pt * fit_scale if role else 10.5 * fit_scale
        if target_height > 0:
            formula_size = min(formula_size, max(target_height / 1.35, 0.5))
        formula_size = self._fit_formula_width(latex, formula_size, container_width * 0.90)
        if not number:
            if not latex:
                self._write_image(container, element, fit_scale, container_width)
                return
            paragraph = container.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(
                float(element.payload.get("space_before_pt", 0.0)) * fit_scale
            )
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(max(formula_size * 1.6, 6.0))
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
            if append_latex_omml(paragraph, latex, formula_size):
                return
            paragraph.add_run(latex)
            return
        table = container.add_table(rows=1, cols=2)
        if target_height > 0:
            table.rows[0].height = Pt(target_height)
            table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        equation_ratio = 0.88
        source_bbox = element.payload.get("source_bbox") or element.payload.get("primary_bbox")
        number_bbox = element.payload.get("number_bbox")
        if source_bbox and number_bbox and len(source_bbox) >= 4 and len(number_bbox) >= 4:
            source_width = max(float(source_bbox[2]) - float(source_bbox[0]), 1.0)
            equation_ratio = min(
                max((float(number_bbox[0]) - float(source_bbox[0])) / source_width, 0.70),
                0.94,
            )
        table_width = min(container_width, target_width) if target_width > 0 else container_width
        number_width = max(table_width * (1.0 - equation_ratio), 24.0)
        widths = (max(table_width - number_width, 1.0), number_width)
        self._format_layout_table(table, widths)
        body = table.cell(0, 0)
        paragraph = body.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        formula_size = self._fit_formula_width(latex, formula_size, widths[0] * 0.94)
        paragraph.paragraph_format.line_spacing = Pt(max(formula_size * 1.6, 6.0))
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        if latex and append_latex_omml(paragraph, latex, formula_size):
            pass
        elif latex:
            paragraph.add_run(latex)
        else:
            data = self._decode_image(element.payload.get("image_base64"))
            if data:
                width = widths[0] * float(element.payload.get("width_fraction", 1.0)) * fit_scale
                paragraph.add_run().add_picture(io.BytesIO(data), width=Pt(max(width, 0.5)))
        number_cell = table.cell(0, 1)
        number_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = number_cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = False
        self._style_run(paragraph.add_run(number), role, fit_scale)
        self._collapse_trailing_paragraph(container)

    def _write_native_table(self, container, element, roles, fit_scale: float, container_width: float) -> None:
        html = element.payload.get("html")
        soup = BeautifulSoup(str(html or "<table><tr><td></td></tr></table>"), "html.parser")
        source = soup.find("table")
        if source is None:
            source = BeautifulSoup("<table><tr><td></td></tr></table>", "html.parser").find("table")
        rows, columns = get_table_dimensions(source)
        rows, columns = max(rows, 1), max(columns, 1)
        table = container.add_table(rows=rows, cols=columns)
        native_marker = OxmlElement("w:tblCaption")
        native_marker.set(qn("w:val"), "docflow-native-table")
        table._tbl.tblPr.append(native_marker)
        rule_style = element.payload.get("table_rule_style") or "grid"
        if rule_style == "grid":
            table.style = "Table Grid"
        elif rule_style == "horizontal":
            head = source.find("thead")
            header_rows = len(head.find_all("tr", recursive=False)) if head else 0
            if not header_rows:
                rows_with_cells = get_table_rows(source)
                header_rows = next(
                    (index for index, row in enumerate(rows_with_cells) if row.find("th", recursive=False) is None),
                    len(rows_with_cells),
                )
            set_horizontal_table_borders(table, header_rows)
        else:
            clear_table_borders(table)
        left_indent = float(element.payload.get("left_indent_pt", 0.0))
        if left_indent > 4.0:
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table_indent = OxmlElement("w:tblInd")
            table_indent.set(qn("w:w"), str(round(left_indent * 20)))
            table_indent.set(qn("w:type"), "dxa")
            table._tbl.tblPr.append(table_indent)
        elif float(element.payload.get("width_fraction", 1.0)) >= 0.90:
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
        else:
            table.alignment = {
                "left": WD_TABLE_ALIGNMENT.LEFT,
                "right": WD_TABLE_ALIGNMENT.RIGHT,
            }.get(element.payload.get("alignment"), WD_TABLE_ALIGNMENT.CENTER)
        column_weights = get_table_column_weights(source, element.payload.get("table_column_width_ratios") or ())
        width_fraction = float(element.payload.get("width_fraction", 1.0))
        if width_fraction >= 0.85:
            width_fraction = 1.0
        table_width = container_width * width_fraction
        column_widths = [table_width * weight / sum(column_weights) for weight in column_weights]
        set_table_col_widths(table, column_widths)
        placements = get_table_cell_placements(source)
        role = roles.get(element.role_id) or self._body_role(roles)
        cell_geometry = {
            (int(row), int(column)): (vertical, top, bottom, start, end)
            for row, column, vertical, top, bottom, start, end in element.payload.get("table_cell_geometry", ())
        }
        base_size = float(element.payload.get("table_font_size_pt") or (role.font_size_pt if role else 10.5))
        minimum_size = float(element.payload.get("table_min_font_size_pt", 0.5))
        font_size = max(round(base_size * fit_scale * 2) / 2.0, minimum_size)
        font_size = fit_table_font_size(
            source,
            font_size,
            table_width,
            column_weights,
            cell_geometry,
            minimum_size,
        )
        source_height = float(element.payload.get("table_height_pt") or 0.0) * fit_scale
        row_ratios = element.payload.get("table_row_height_ratios") or ()
        if len(row_ratios) == rows:
            source_row_heights = [source_height * float(ratio) for ratio in row_ratios]
        else:
            source_row_height = 0.0 if element.payload.get("table_content_fit") else source_height / rows
            source_row_heights = [source_row_height] * rows
        row_heights = [max(source_row_heights[row], font_size * 1.2 + 2.0) for row in range(rows)]
        required_rows = [font_size * 1.2 + 2.0 for _ in range(rows)]
        visual_rows = [False for _ in range(rows)]
        has_vertical_merge = any(row_span > 1 for _row, _column, row_span, _span, _cell in placements)
        for row_index, column, row_span, span, cell in placements:
            cell_width = sum(column_widths[column : column + span])
            geometry = cell_geometry.get((row_index, column), ("center", 0, 0, 0, 0))
            available_width = table_cell_available_width(cell_width, geometry[3], geometry[4])
            lines = max(1, estimate_table_cell_lines(cell, font_size, available_width))
            required_per_row = (lines * font_size * 1.2 + 2.0) / max(row_span, 1)
            for target_row in range(row_index, min(row_index + row_span, rows)):
                required_rows[target_row] = max(required_rows[target_row], required_per_row)
                visual_rows[target_row] = visual_rows[target_row] or bool(cell.find("img"))
                row_heights[target_row] = max(row_heights[target_row], required_per_row)
        page_height = float(element.payload.get("page_height_px") or 0.0) * float(
            element.payload.get("source_scale") or 1.0
        )
        compact_text_rows = bool(element.payload.get("table_auto_height")) or (
            page_height > 0.0 and source_height > page_height * 0.45
        )
        if compact_text_rows:
            for row in range(rows):
                row_heights[row] = (
                    max(row_heights[row], required_rows[row])
                    if visual_rows[row]
                    else required_rows[row]
                )
        for row_index, (row, row_height) in enumerate(zip(table.rows, row_heights)):
            row.height = Pt(row_height)
            row.height_rule = (
                WD_ROW_HEIGHT_RULE.AT_LEAST
                if element.payload.get("table_auto_height") or has_vertical_merge
                else WD_ROW_HEIGHT_RULE.EXACTLY
            )
        row_styles = {
            int(row): (fill, text_color)
            for row, fill, text_color in element.payload.get("table_row_styles", ())
        }
        cell_styles = {
            (int(row), int(column)): (alignment, bool(bold))
            for row, column, alignment, bold in element.payload.get("table_cell_styles", ())
        }
        for row_index, column_index, row_span, column_span, cell_source in placements:
            cell = table.cell(row_index, column_index)
            cell_width = sum(column_widths[column_index : column_index + column_span])
            cell_height = sum(row_heights[row_index : row_index + row_span])
            if row_span > 1 or column_span > 1:
                cell = cell.merge(table.cell(row_index + row_span - 1, column_index + column_span - 1))
            alignment, bold = cell_styles.get(
                (row_index, column_index),
                ("center" if cell_source.name == "th" else "left", cell_source.name == "th"),
            )
            vertical, top_ratio, bottom_ratio, start_ratio, end_ratio = cell_geometry.get(
                (row_index, column_index),
                ("center", 0.0, 0.0, 0.0, 0.0),
            )
            start_ratio, end_ratio = table_cell_padding_ratios(start_ratio, end_ratio)
            available_width = table_cell_available_width(cell_width, start_ratio, end_ratio)
            cell_text = "\n".join(get_table_cell_text_lines(cell_source))
            text_lines = (
                max(1, estimate_table_cell_lines(cell_source, font_size, available_width))
                if cell_text
                else 0
            )
            cell.vertical_alignment = {
                "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
            }.get(vertical, WD_CELL_VERTICAL_ALIGNMENT.CENTER)
            vertical_padding_height = cell_height / max(row_span, 1)
            top_padding = top_ratio * vertical_padding_height
            bottom_padding = bottom_ratio * vertical_padding_height
            if text_lines:
                max_vertical_padding = max(cell_height - text_lines * font_size * 1.2 - 2.0, 0.0)
                if top_padding + bottom_padding > max_vertical_padding:
                    padding_scale = max_vertical_padding / max(top_padding + bottom_padding, 1e-9)
                    top_padding *= padding_scale
                    bottom_padding *= padding_scale
            set_cell_margins(
                cell,
                top=round(top_padding * 20),
                bottom=round(bottom_padding * 20),
                start=round(start_ratio * cell_width * 20),
                end=round(end_ratio * cell_width * 20),
            )
            image_nodes = cell_source.find_all("img", src=True)[:4]
            formula_nodes = cell_source.find_all("span", class_="docflow-formula")
            self._clear_container(cell)
            annotated = has_table_cell_paragraphs(cell_source)
            content_nodes = get_table_cell_content_nodes(cell_source) if annotated else ()
            cell_height = sum(row_heights[row_index : row_index + row_span])
            image_height = max(
                (
                    cell_height
                    - max(text_lines * font_size * 1.2, font_size * 3.0 if cell_text else 0.0)
                    - 4.0
                )
                / max(len(image_nodes), 1),
                8.0,
            )
            if annotated:
                for content_node in content_nodes:
                    if content_node.name == "p":
                        try:
                            space_before = float(content_node.get("data-space-before") or 0.0)
                        except (TypeError, ValueError):
                            space_before = 0.0
                        self._add_table_text_paragraph(
                            cell,
                            content_node.get_text("\n", strip=True),
                            str(content_node.get("data-align") or alignment),
                            content_node.get("data-bold") == "true" or bold,
                            space_before,
                            font_size,
                            role,
                            element,
                            row_index,
                            row_styles,
                        )
                        continue
                    if content_node.name == "img":
                        self._add_table_image_paragraph(cell, content_node, available_width, image_height)
                        continue
                    self._add_table_formula_paragraph(cell, content_node, font_size, cell_height, cell_width)
                if not content_nodes:
                    self._add_table_text_paragraph(
                        cell,
                        "",
                        alignment,
                        bold,
                        0.0,
                        font_size,
                        role,
                        element,
                        row_index,
                        row_styles,
                    )
                if row_index in row_styles:
                    fill, _text_color = row_styles[row_index]
                    self._shade(cell._tc.get_or_add_tcPr(), fill)
                continue
            for image_index, image_node in enumerate(image_nodes):
                self._add_table_image_paragraph(cell, image_node, available_width, image_height)
            for formula_index, formula_node in enumerate(formula_nodes):
                self._add_table_formula_paragraph(cell, formula_node, font_size, cell_height, cell_width)
            if cell_text:
                self._add_table_text_paragraph(
                    cell,
                    cell_text,
                    alignment,
                    bold,
                    0.0,
                    font_size,
                    role,
                    element,
                    row_index,
                    row_styles,
                )
            elif not image_nodes and not formula_nodes:
                self._add_table_text_paragraph(
                    cell,
                    "",
                    alignment,
                    bold,
                    0.0,
                    font_size,
                    role,
                    element,
                    row_index,
                    row_styles,
                )
            if row_index in row_styles:
                fill, _text_color = row_styles[row_index]
                self._shade(cell._tc.get_or_add_tcPr(), fill)
        self._collapse_trailing_paragraph(container)

    def _add_table_text_paragraph(
        self,
        cell,
        text: str,
        alignment: str,
        bold: bool,
        space_before: float,
        font_size: float,
        role,
        element,
        row_index: int,
        row_styles,
    ):
        paragraph = cell.add_paragraph()
        paragraph.alignment = _ALIGNMENT.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)
        paragraph.paragraph_format.space_before = Pt(max(space_before, 0.0) * font_size)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = Pt(font_size * 1.2)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        paragraph.paragraph_format.widow_control = False
        paragraph.paragraph_format.keep_together = False
        paragraph.paragraph_format.keep_with_next = False
        lines = str(text).splitlines() or [""]
        run = paragraph.add_run(lines[0])
        self._style_run(
            run,
            role,
            1.0,
            font_size_pt=font_size,
            font_family=element.payload.get("font_family"),
        )
        run.bold = bold
        for line in lines[1:]:
            run.add_break()
            run.add_text(line)
        if row_index in row_styles:
            _fill, text_color = row_styles[row_index]
            run.font.color.rgb = RGBColor.from_string(text_color.lstrip("#"))
        return paragraph

    def _add_table_image_paragraph(self, cell, image_node, cell_width: float, image_height: float) -> None:
        data = self._decode_image(image_node.get("src"))
        if not data:
            return
        paragraph = cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        image_width, fitted_height = self._table_image_size(data, cell_width, image_height)
        paragraph.add_run().add_picture(io.BytesIO(data), width=Pt(image_width), height=Pt(fitted_height))

    @staticmethod
    def _add_table_formula_paragraph(cell, formula_node, font_size: float, cell_height: float, available_width: float) -> None:
        latex = str(formula_node.get("data-latex") or "").strip()
        paragraph = cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        formula_size = min(font_size, max(cell_height / 1.35, 0.5))
        formula_size = ReflowDocxRenderer._fit_formula_width(latex, formula_size, available_width)
        paragraph.paragraph_format.line_spacing = Pt(max(formula_size * 1.6, 6.0))
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        if not latex or append_latex_omml(paragraph, latex, formula_size):
            return
        paragraph.add_run(f"${latex}$")

    @staticmethod
    def _fit_formula_width(latex: str, formula_size: float, available_width: float) -> float:
        if not latex:
            return formula_size
        try:
            pdf = MathJaxFormulaRenderer.render_pdf(latex)
            with fitz.open(stream=pdf, filetype="pdf") as document:
                page = document[0]
                drawings = [drawing["rect"] for drawing in page.get_drawings()]
                content = drawings[0]
                for rect in drawings[1:]:
                    content |= rect
            if content.width <= 0:
                return formula_size
            fitted = formula_size * available_width * 0.94 / content.width
            return max(min(formula_size, fitted), 5.0)
        except Exception:
            return formula_size

    @staticmethod
    def _table_image_size(data: bytes, cell_width: float, cell_height: float = 47.5) -> tuple[float, float]:
        return ReflowDocxRenderer._image_size(
            data,
            max(8.0, cell_width - 4.0),
            max(8.0, cell_height - 4.0),
        )

    @staticmethod
    def _image_size(data: bytes, max_width: float, max_height: float) -> tuple[float, float]:
        try:
            with Image.open(io.BytesIO(data)) as image:
                width, height = image.size
            scale = min(max_width / max(width, 1), max_height / max(height, 1))
            return max(1.0, width * scale), max(1.0, height * scale)
        except (OSError, ValueError):
            return max_width, max_height if math.isfinite(max_height) else max_width

    @staticmethod
    def _format_layout_table(table, widths, gutter_pt: float = 0.0) -> None:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        clear_table_borders(table)
        column_count = len(widths)
        layout_widths = tuple(
            width + gutter_pt * ((column > 0) + (column < column_count - 1)) / 2.0
            for column, width in enumerate(widths)
        )
        set_table_col_widths(table, layout_widths)
        for row in table.rows:
            for column, cell in enumerate(row.cells):
                half_gutter_twips = int(max(gutter_pt, 0.0) * 10)
                set_cell_margins(
                    cell,
                    top=0,
                    bottom=0,
                    start=half_gutter_twips if column > 0 else 0,
                    end=half_gutter_twips if column < len(row.cells) - 1 else 0,
                )

    @staticmethod
    def _style_run(run, role, fit_scale: float, font_size_pt=None, font_family=None) -> None:
        if role is None:
            if font_size_pt is None:
                return
            run.font.name = "Times New Roman"
            run.font.size = Pt(max(round(float(font_size_pt) * fit_scale * 2) / 2.0, 0.5))
            fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                run._element.get_or_add_rPr().insert(0, fonts)
            fonts.set(qn("w:eastAsia"), font_family or "宋体")
            return
        run.font.name = role.western_font_family
        size = float(font_size_pt) if font_size_pt is not None else role.font_size_pt
        run.font.size = Pt(max(round(size * fit_scale * 2) / 2.0, 0.5))
        run.bold = role.bold
        run.italic = role.italic
        run.font.color.rgb = RGBColor.from_string(role.color.lstrip("#"))
        fonts = run._element.get_or_add_rPr().find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            run._element.get_or_add_rPr().insert(0, fonts)
        fonts.set(qn("w:eastAsia"), font_family or role.font_family)

    @staticmethod
    def _shade(properties, color) -> None:
        if not color:
            return
        shading = properties.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            properties.append(shading)
        shading.set(qn("w:fill"), str(color).lstrip("#"))

    @staticmethod
    def _body_role(roles):
        return next((role for role_id, role in roles.items() if role_id.startswith("body_")), next(iter(roles.values()), None))

    @staticmethod
    def _apply_payload_text_color(paragraph, element) -> None:
        color = str(element.payload.get("text_color") or "").strip().lstrip("#")
        if len(color) != 6:
            return
        for run in paragraph.runs:
            run.font.color.rgb = RGBColor.from_string(color)

    def _render_header_texts(self, container, elements, roles, fit_scale: float, usable_width: float) -> None:
        role = self._body_role(roles)
        for element in elements:
            rows = element.text_rows
            if len(rows) == 1 and len(rows[0].spans) > 1:
                self._write_header_span_table(container, element, rows[0].spans, role, fit_scale, usable_width)
                continue
            paragraph = container.add_paragraph()
            self._write_text(paragraph, element, {role.role_id: role} if role else {}, fit_scale, usable_width)
            self._apply_payload_text_color(paragraph, element)
        if len(elements) > 1:
            color = str(elements[0].payload.get("text_color") or "").strip().lstrip("#")
            if len(color) == 6:
                channels = [int(color[index : index + 2], 16) for index in (0, 2, 4)]
                if max(channels) - min(channels) >= 20:
                    rule = container.add_paragraph()
                    rule.paragraph_format.space_before = Pt(0)
                    rule.paragraph_format.space_after = Pt(0)
                    rule.paragraph_format.line_spacing = Pt(1)
                    rule.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    properties = rule._p.get_or_add_pPr()
                    borders = OxmlElement("w:pBdr")
                    bottom = OxmlElement("w:bottom")
                    bottom.set(qn("w:val"), "single")
                    bottom.set(qn("w:sz"), "8")
                    bottom.set(qn("w:space"), "1")
                    bottom.set(qn("w:color"), color)
                    borders.append(bottom)
                    properties.append(borders)

    def _write_header_span_table(self, container, element, spans, role, fit_scale: float, usable_width: float) -> None:
        source_bbox = tuple(float(value) for value in element.payload.get("source_bbox") or (0, 0, 1, 1))
        source_width = max(source_bbox[2] - source_bbox[0], 1.0)
        split_x = source_bbox[0] + source_width * 0.72
        left_spans = sorted(
            (span for span in spans if (span.bbox.x1 + span.bbox.x2) / 2.0 < split_x),
            key=lambda span: span.bbox.y1,
        )
        right_spans = sorted(
            (span for span in spans if (span.bbox.x1 + span.bbox.x2) / 2.0 >= split_x),
            key=lambda span: span.bbox.y1,
        )
        if not left_spans or not right_spans:
            self._render_header_texts(
                container,
                (replace(element, text_rows=()),),
                {role.role_id: role} if role else {},
                fit_scale,
                usable_width,
            )
            return
        rows = max(len(left_spans), len(right_spans))
        table = container.add_table(rows=rows, cols=2, width=Pt(usable_width))
        left_width = max(span.bbox.x2 for span in left_spans) - min(span.bbox.x1 for span in left_spans)
        right_width = max(span.bbox.x2 for span in right_spans) - min(span.bbox.x1 for span in right_spans)
        ratios = (left_width / max(left_width + right_width, 1.0), right_width / max(left_width + right_width, 1.0))
        self._format_layout_table(table, tuple(usable_width * ratio for ratio in ratios))
        merged_right = None
        if len(right_spans) == 1 and rows > 1:
            merged_right = table.cell(0, 1).merge(table.cell(rows - 1, 1))
        for row in table.rows:
            for cell in row.cells:
                self._clear_container(cell)
        if len(right_spans) > 1:
            target = table.cell(0, 1)
            for span, row_index in zip(right_spans, range(rows)):
                self._write_header_span(table.cell(row_index, 1), element, span, role, fit_scale)
        else:
            target = merged_right or table.cell(0, 1)
            target.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            self._write_header_span(target, element, right_spans[0], role, fit_scale)
        for span, row_index in zip(left_spans, range(rows)):
            self._write_header_span(table.cell(row_index, 0), element, span, role, fit_scale)

    @staticmethod
    def _header_span_size(element, row) -> float:
        source_scale = float(element.payload.get("source_scale", 1.0))
        return max(float(row.ink_height_px or row.bbox.height) * source_scale * 0.90, 4.0)

    def _write_header_span(self, cell, element, span, role, fit_scale: float) -> None:
        paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        char_count = max(len(span.text), 1)
        glyph_extent = min(float(span.bbox.height), float(span.bbox.width) / char_count)
        size = max(glyph_extent * float(element.payload.get("source_scale", 1.0)) * 0.90, 4.0)
        paragraph.paragraph_format.line_spacing = Pt(size * 1.05)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        run = paragraph.add_run(span.text)
        self._style_run(run, role, 1.0, font_size_pt=size, font_family=element.payload.get("font_family"))
        self._apply_payload_text_color(paragraph, element)

    def _render_footer_texts(self, container, elements, roles, fit_scale: float, usable_width: float) -> None:
        role = self._body_role(roles)
        items = []
        for element in elements:
            rows = element.text_rows
            if element.kind == "footer" and rows and rows[0].spans:
                items.extend((element, span.text, span.bbox) for span in rows[0].spans)
            else:
                bbox = Rect.from_sequence(element.payload.get("source_bbox") or (0, 0, 1, 1))
                items.append((element, element.text, bbox))
        if not items:
            return
        if len(items) == 1:
            element, text, _bbox = items[0]
            paragraph = container.add_paragraph()
            self._write_text(paragraph, replace(element, text=text), {role.role_id: role} if role else {}, fit_scale, usable_width)
            self._apply_payload_text_color(paragraph, element)
            return
        table = container.add_table(rows=1, cols=len(items), width=Pt(usable_width))
        widths = [
            max(bbox.width, 1.0)
            for _element, _text, bbox in items
        ]
        total_width = sum(widths)
        self._format_layout_table(
            table,
            tuple(usable_width * width / total_width for width in widths),
        )
        if any(element.kind == "footer" for element, _text, _bbox in items):
            frame_left = min(bbox.x1 for _element, _text, bbox in items)
            frame_right = max(bbox.x2 for _element, _text, bbox in items)
        else:
            page_width = max(float(elements[0].payload.get("page_width_px", 1.0)), 1.0)
            frame_left, frame_right = 0.0, page_width
        for (element, text, bbox), cell in zip(items, table.rows[0].cells):
            self._clear_container(cell)
            paragraph = cell.add_paragraph()
            paragraph.alignment = _ALIGNMENT[
                self._infer_furniture_alignment(bbox, frame_left, frame_right)
            ]
            source_scale = float(element.payload.get("source_scale", 1.0))
            size = max(bbox.height * source_scale * 0.90, 4.0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(size * 1.05)
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            if self._is_page_number_element(element):
                layouts = element.text_layout or (
                    resolve_text_layout(element, role, max(bbox.width, 1.0), fit_scale) if role is not None else ()
                )
                layout = layouts[0] if layouts else TextParagraphLayout(
                    str(text),
                    self._infer_furniture_alignment(bbox, frame_left, frame_right),
                    size,
                    size * 1.05,
                )
                layout = self._page_number_layout(element, layout)
                self._append_page_number_runs(paragraph, replace(element, text=text), role, layout)
            else:
                run = paragraph.add_run(text)
                self._style_run(run, role, 1.0, font_size_pt=size, font_family=element.payload.get("font_family"))
            self._apply_payload_text_color(paragraph, element)

    @staticmethod
    def _infer_furniture_alignment(bbox, frame_left: float, frame_right: float) -> str:
        width = max(frame_right - frame_left, 1.0)
        relative = ((bbox.x1 + bbox.x2) / 2.0 - frame_left) / width
        if abs(relative - 0.5) <= 0.17:
            return "center"
        return "left" if relative < 0.5 else "right"

    @staticmethod
    def _visual_text(element) -> str:
        return visual_text(element)

    @staticmethod
    def _content_bbox(element) -> Rect:
        if element.content_bbox is not None:
            return element.content_bbox
        return Rect.from_sequence(element.payload.get("source_bbox") or (0, 0, 1, 1))

    @staticmethod
    def _decode_image(value):
        try:
            if isinstance(value, str) and value.startswith("data:"):
                value = value.split(",", 1)[1]
            return base64.b64decode(value) if value else None
        except (IndexError, ValueError, TypeError):
            return None

    @staticmethod
    def _usable_width(geometry) -> float:
        return geometry.width_pt - geometry.margin_left_pt - geometry.margin_right_pt

    @staticmethod
    def _clear_container(container) -> None:
        for paragraph in container.paragraphs:
            paragraph._element.getparent().remove(paragraph._element)

    @staticmethod
    def _collapse_trailing_paragraph(container) -> None:
        if not container.paragraphs or container.paragraphs[-1].text:
            return
        paragraph = container.paragraphs[-1]
        following = paragraph._p.getnext()
        if following is not None and following.tag != qn("w:sectPr"):
            return
        ReflowDocxRenderer._collapse_section_break(paragraph)
